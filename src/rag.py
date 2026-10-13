"""PDF RAG utilities and Nepali Constitution article loader."""

from __future__ import annotations

import hashlib
import io
import re
import uuid

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import numpy as np

from docx import Document as WordDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS

# --------------------------------------------------
# 1. Local Embedding Model
# --------------------------------------------------


@lru_cache(maxsize=1)
def _encoder():
    # Load MiniLM once and reuse it for embeddings.
    from sentence_transformers import SentenceTransformer

    return SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")


class LocalEmbeddings(Embeddings):
    """Connect SentenceTransformers with LangChain."""

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return _encoder().encode(texts, normalize_embeddings=True).tolist()

    def embed_query(self, text: str) -> list[float]:
        return _encoder().encode([text], normalize_embeddings=True)[0].tolist()


# --------------------------------------------------
# 2. RAG Index Data Structure
# --------------------------------------------------


@dataclass
class RAGIndex:
    store: Chroma
    chunks: list[Document]
    document_id: str
    filename: str
    total_pages: int
    embeddings: LocalEmbeddings
    vectors: np.ndarray
    lexical: TfidfVectorizer
    lexical_matrix: object


# --------------------------------------------------
# 3. Existing PDF Ingestion
# --------------------------------------------------


def pdf_to_chunks(pdf_bytes: bytes, filename: str) -> tuple[int, list[Document]]:
    """Extract PDF text and preserve page metadata."""

    reader = PdfReader(io.BytesIO(pdf_bytes))
    page_docs = []

    doc_id = hashlib.sha256(pdf_bytes).hexdigest()

    for page_number, page in enumerate(reader.pages, 1):
        content = (page.extract_text() or "").strip()

        if content:
            page_docs.append(
                Document(
                    page_content=content,
                    metadata={
                        "source": filename,
                        "page": page_number,
                        "document_id": doc_id,
                    },
                )
            )

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)

    chunks = splitter.split_documents(page_docs)

    # Add stable IDs for citations and debugging.
    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = f"{doc_id[:10]}-{i:04d}"

    return len(reader.pages), chunks


# --------------------------------------------------
# 4. Existing Vector Database Indexing
# --------------------------------------------------


def build_index(pdf_bytes: bytes, filename: str) -> RAGIndex:
    """Create a Chroma vector index and TF-IDF index."""

    # Use article-aware indexing only for the verified
    # English Constitution PDF. Preserve generic PDF mode.
    constitution_path = Path("data/constitution/constitution_en.pdf")

    use_article_index = (
        filename == "constitution_en.pdf"
        and constitution_path.is_file()
        and constitution_path.read_bytes() == pdf_bytes
    )

    if use_article_index:
        articles = load_english_constitution_articles(
            constitution_path
        )

        # The loader validates this specific 202-page PDF.
        pages = 202

        article_splitter = RecursiveCharacterTextSplitter(
            chunk_size=2000,
            chunk_overlap=200,
        )

        chunks = article_splitter.split_documents(articles)

        # Preserve article identity in every indexed chunk.
        for i, chunk in enumerate(chunks):
            chunk.metadata["chunk_id"] = (
                f"{chunk.metadata['document_id'][:10]}"
                f"-a{chunk.metadata['article']:03d}-{i:04d}"
            )
    else:
        pages, chunks = pdf_to_chunks(pdf_bytes, filename)

    if not chunks:
        raise ValueError("PDF has no extractable text. " "Scanned PDFs require OCR.")

    embedding = LocalEmbeddings()

    vectors = np.asarray(embedding.embed_documents([d.page_content for d in chunks]))

    # Keyword-based search.
    lexical = TfidfVectorizer(ngram_range=(1, 2), token_pattern=r"(?u)\b\w+\b")

    lexical_matrix = lexical.fit_transform([d.page_content for d in chunks])

    # Create an isolated in-memory Chroma collection.
    store = Chroma.from_documents(
        documents=chunks,
        embedding=embedding,
        collection_name=f"ragtrust_{uuid.uuid4().hex}",
    )

    return RAGIndex(
        store=store,
        chunks=chunks,
        document_id=chunks[0].metadata["document_id"],
        filename=filename,
        total_pages=pages,
        embeddings=embedding,
        vectors=vectors,
        lexical=lexical,
        lexical_matrix=lexical_matrix,
    )


# --------------------------------------------------
# 5. Existing Hybrid Retrieval
# --------------------------------------------------



def extract_article_references(question: str) -> list[int]:
    """Extract explicitly referenced article numbers and short lists."""

    lead_pattern = re.compile(
        r"(?:\barticles?\b|धारा)\s*([0-9०-९]+)(?![0-9०-९])",
        flags=re.IGNORECASE,
    )

    continuation_pattern = re.compile(
        r"\s*(?:,\s*(?:(?:and|or)\b\s*|र\s*)?|"
        r"(?:and|or)\b\s+|र\s+)"
        r"([0-9०-९]+)(?![0-9०-९])",
        flags=re.IGNORECASE,
    )

    digit_map = str.maketrans("०१२३४५६७८९", "0123456789")
    numbers = []

    for lead in lead_pattern.finditer(question):
        numbers.append(int(lead.group(1).translate(digit_map)))

        position = lead.end()

        while True:
            continuation = continuation_pattern.match(
                question, position
            )

            if continuation is None:
                break

            numbers.append(
                int(continuation.group(1).translate(digit_map))
            )

            position = continuation.end()

    # Remove duplicates while keeping their original order.
    return list(dict.fromkeys(numbers))


def hybrid_retrieve(
    index: RAGIndex, question: str, k: int = 5
) -> tuple[list[Document], float]:
    """Combine dense retrieval and TF-IDF using RRF."""

    count = min(len(index.chunks), 12)

    # Dense vector search.
    dense_docs = index.store.similarity_search(
        question, k=count, filter={"document_id": index.document_id}
    )

    dense_ids = [d.metadata["chunk_id"] for d in dense_docs]

    # Lexical keyword search.
    lexical_scores = index.lexical.transform([question]) @ index.lexical_matrix.T

    lexical_scores = np.asarray(lexical_scores.toarray()).ravel()

    lexical_positions = np.argsort(-lexical_scores)[:count]

    lexical_ids = [
        index.chunks[i].metadata["chunk_id"]
        for i in lexical_positions
        if lexical_scores[i] > 0
    ]

    # Reciprocal Rank Fusion.
    rank_scores: dict[str, float] = {}

    for ranking in (dense_ids, lexical_ids):
        for rank, chunk_id in enumerate(ranking, 1):
            rank_scores[chunk_id] = rank_scores.get(chunk_id, 0.0) + 1 / (60 + rank)

    by_id = {d.metadata["chunk_id"]: d for d in index.chunks}

    # Rank results from existing dense + lexical retrieval.
    ranked_ids = sorted(
        rank_scores,
        key=rank_scores.get,
        reverse=True
    )

    # Detect explicit article references in English or Nepali.
    # Examples: Article 18, धारा १८
    # Recognize single articles and article-number lists.
    article_numbers = extract_article_references(question)
    article_match = bool(article_numbers)

    # Selective Constitution article-title reranking.
    # Apply only when the query has no explicit article number.
    if (
        not article_match
        and index.filename == "constitution_en.pdf"
        and ranked_ids
    ):
        article_starts = {}

        for doc in index.chunks:
            number = doc.metadata.get("article")

            if isinstance(number, int) and 1 <= number <= 308:
                article_starts.setdefault(number, doc)

        # Require the complete article-aware index.
        if len(article_starts) == 308:
            titles = [
                extract_english_article_title(article_starts[n])
                for n in range(1, 309)
            ]

            title_stopwords = set(ENGLISH_STOP_WORDS) | {
                "constitution", "constitutional", "nepal",
                "nepali", "article", "articles",
                "right", "rights", "law",
                "provision", "provisions",
            }

            title_vectorizer = TfidfVectorizer(
                stop_words=sorted(title_stopwords),
                ngram_range=(1, 2),
            )

            title_matrix = title_vectorizer.fit_transform(titles)

            title_scores = np.asarray(
                (
                    title_vectorizer.transform([question])
                    @ title_matrix.T
                ).toarray()
            ).ravel()

            def title_adjusted_score(item):
                rank, chunk_id = item
                number = by_id[chunk_id].metadata.get("article")

                title_score = (
                    float(title_scores[number - 1])
                    if isinstance(number, int)
                    and 1 <= number <= 308
                    else 0.0
                )

                boost = (
                    0.025 * title_score
                    if title_score > 0.12
                    else 0.0
                )

                return 1 / (60 + rank) + boost

            ranked_ids = [
                chunk_id
                for _, chunk_id in sorted(
                    enumerate(ranked_ids, 1),
                    key=title_adjusted_score,
                    reverse=True,
                )
            ]

    # Multi-article evidence prioritization.
    if article_match:
        is_english_constitution = (
            index.filename == "constitution_en.pdf"
        )

        if is_english_constitution:
            ranked_ids = [
                cid for cid in ranked_ids
                if 1 <= by_id[cid].metadata.get("page", 0) <= 184
            ]

        matches_by_article = []

        for article_no in article_numbers:
            heading = re.compile(
                rf"(?m)^[ \t]*{article_no}[.।][ \t]+\S"
            )

            matches_for_article = [
                doc.metadata["chunk_id"]
                for doc in index.chunks
                if (
                    not is_english_constitution
                    or 1 <= doc.metadata.get("page", 0) <= 184
                )
                and (
                    doc.metadata.get("article") == article_no
                    or heading.search(doc.page_content)
                )
            ]

            matches_by_article.append(matches_for_article)

        if len(article_numbers) > 1:
            # Reserve one evidence chunk for each requested article.
            primary = [
                ids[0] for ids in matches_by_article if ids
            ]

            # Keep remaining article chunks available afterward.
            additional = [
                cid
                for ids in matches_by_article
                for cid in ids[1:]
            ]

            ranked_ids = list(dict.fromkeys(
                primary + additional + ranked_ids
            ))
        else:
            exact_ids = matches_by_article[0]
            ranked_ids = list(dict.fromkeys(
                exact_ids + ranked_ids
            ))

    chosen_ids = ranked_ids[:k]

    chosen = [by_id[cid] for cid in chosen_ids]

    # Calculate cosine similarity.
    question_vector = np.asarray(index.embeddings.embed_query(question))

    selected_indices = [index.chunks.index(d) for d in chosen]

    best_similarity = max(
        (float(index.vectors[i] @ question_vector) for i in selected_indices),
        default=0.0,
    )

    return chosen, best_similarity


# --------------------------------------------------
# 6. Nepali Constitution Configuration
# --------------------------------------------------

# Convert Nepali digits to English digits.
NEPALI_DIGITS = str.maketrans("०१२३४५६७८९", "0123456789")

# Recognise article headings with or without a dot.
#
# Example:
# १३४.      मुद्दा सार्न सक्नेः
# १३५       बहस पैरवी गर्न नपाउनेः

ARTICLE_PATTERN = re.compile(r"^([०-९0-9]{1,3})(?:[.।][ \t]*|\t+| {2,})(\S.*)$")

# Path to the original Nepali Constitution DOCX.
CONSTITUTION_PATH = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "constitution"
    / "नेपालको संविधान (जारी हुँदाको).docx"
)


# --------------------------------------------------
# 7. Constitution Article Loader
# --------------------------------------------------


def load_constitution_articles(path: str | Path = CONSTITUTION_PATH) -> list[Document]:
    """
    Extract the original 2072 Constitution by article.

    Every article becomes one LangChain Document.

    Important:
    This is the original Constitution, not a
    consolidated version with later amendments.
    Schedules and Word tables are not processed here.
    """

    path = Path(path)

    word_doc = WordDocument(path)

    # Unique ID for this exact source file.
    document_id = hashlib.sha256(path.read_bytes()).hexdigest()

    articles: list[Document] = []

    current_number: int | None = None
    current_text: list[str] = []

    def save_article():
        """Save the current article with metadata."""

        if current_number is None:
            return

        articles.append(
            Document(
                page_content="\n".join(current_text),
                metadata={
                    "source": path.name,
                    "document_id": document_id,
                    "chunk_id": (f"constitution-2072-article-" f"{current_number:03d}"),
                    "language": "ne",
                    "article": current_number,
                    "version": "2072_original",
                    "status": "historical_original",
                },
            )
        )

    # Process each Word paragraph.
    for paragraph in word_doc.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        # Stop when schedules begin after Article 308.
        if current_number == 308 and re.match(r"^अनुसूची\s*[-–—]?\s*[०-९0-9]+", text):
            break

        match = ARTICLE_PATTERN.match(text)

        if match:
            number = int(match.group(1).translate(NEPALI_DIGITS))

            if 1 <= number <= 308:

                # Articles must appear in sequence.
                expected = current_number + 1 if current_number is not None else 1

                if number != expected:
                    raise ValueError(
                        f"Unexpected article numbering: "
                        f"expected {expected}, found {number}. "
                        f"Text: {text[:90]}"
                    )

                # Save the previous article.
                save_article()

                # Start the newly detected article.
                current_number = number
                current_text = [text]

                continue

        # Add subclauses and continuation paragraphs.
        if current_number is not None:
            current_text.append(text)

    # Save Article 308.
    save_article()

    # Confirm all 308 articles were extracted.
    numbers = [item.metadata["article"] for item in articles]

    if numbers != list(range(1, 309)):

        missing = sorted(set(range(1, 309)) - set(numbers))

        raise ValueError(f"Article validation failed; missing: {missing}")

    return articles


# --------------------------------------------------
# 8. Constitution Loader Test
# --------------------------------------------------

if __name__ == "__main__":

    articles = load_constitution_articles()

    print("Total articles:", len(articles))

    print("First article:", articles[0].metadata["article"])

    print("Last article:", articles[-1].metadata["article"])

    # Article 135 previously failed due to formatting.
    print("Article 135:", articles[134].page_content[:300])

    print("Version:", articles[134].metadata["version"])



# --------------------------------------------------
# 9. English Constitution Article Loader
# --------------------------------------------------

def load_english_constitution_articles(
    path: str | Path
) -> list[Document]:
    """
    Extract Articles 1-308 from the English Constitution PDF.

    This parser is specific to the verified 202-page PDF:
    - Real articles start on PDF page 3.
    - Article 308 ends on PDF page 184.
    - Schedules start on PDF page 185.

    Each article is kept as one logical Document.
    """

    path = Path(path)
    pdf_bytes = path.read_bytes()
    reader = PdfReader(io.BytesIO(pdf_bytes))

    if len(reader.pages) != 202:
        raise ValueError("Unexpected English Constitution PDF version.")

    document_id = hashlib.sha256(pdf_bytes).hexdigest()

    # Match actual numbered headings within the article section.
    heading_pattern = re.compile(
        r"(?m)^[ \t]*(\d{1,3})\.[ \t]+[^\n]+"
    )

    articles = []
    current_number = None
    current_start_page = None
    current_end_page = None
    current_parts = []

    def save_article():
        if current_number is None:
            return

        articles.append(
            Document(
                page_content="\\n".join(current_parts).strip(),
                metadata={
                    "source": path.name,
                    "document_id": document_id,
                    "chunk_id": f"constitution-en-{current_number:03d}",
                    "language": "en",
                    "article": current_number,
                    "page": current_start_page,
                    "page_end": current_end_page,
                    "version": "amended_through_2020"
                }
            )
        )

    # Skip preamble/amendment list and exclude schedules.
    for page_number in range(3, 185):
        text = reader.pages[page_number - 1].extract_text() or ""

        matches = [
            match
            for match in heading_pattern.finditer(text)
            if 1 <= int(match.group(1)) <= 308
        ]

        cursor = 0

        for match in matches:
            number = int(match.group(1))

            if current_number is not None:
                current_parts.append(text[cursor:match.start()])
                current_end_page = page_number
                save_article()

            expected = len(articles) + 1

            if number != expected:
                raise ValueError(
                    f"Expected Article {expected}, found {number} "
                    f"on PDF page {page_number}"
                )

            current_number = number
            current_start_page = page_number
            current_end_page = page_number
            current_parts = []

            cursor = match.start()

        if current_number is not None:
            current_parts.append(text[cursor:])
            current_end_page = page_number

    save_article()

    numbers = [doc.metadata["article"] for doc in articles]

    if numbers != list(range(1, 309)):
        raise ValueError("Expected exactly 308 sequential articles.")

    return articles



def extract_english_article_title(article: Document) -> str:
    """Extract the heading without including the article body."""

    content = article.page_content

    heading = re.match(
        r"^\s*(\d{1,3})\.\s*",
        content,
    )

    if not heading:
        raise ValueError("Article heading not found.")

    number = int(heading.group(1))

    if number != article.metadata.get("article"):
        raise ValueError(
            f"Article number mismatch: {number}"
        )

    # The heading ends at a colon or the first subsection.
    boundary = re.search(
        r":|\(\s*1\s*\)",
        content[heading.end():],
    )

    if not boundary:
        raise ValueError(
            f"Title boundary not found for Article {number}"
        )

    title = content[
        heading.end():heading.end() + boundary.start()
    ]

    title = " ".join(title.split()).strip()

    if not title:
        raise ValueError(f"Empty title for Article {number}")

    return title
